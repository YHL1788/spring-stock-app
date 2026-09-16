from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\User\spring-stock-app\Candy老师美东9月14日夜盘直播详细总结.docx")
VIDEO_URL = "https://www.youtube.com/watch?v=8KsOlz83EzQ"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, font_name="Microsoft YaHei"):
    run.font.name = font_name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, color="1F5D8F", underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
    r_fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    if underline:
        u_node = OxmlElement("w:u")
        u_node.set(qn("w:val"), "single")
        r_pr.append(u_node)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, 8.5, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, 8.5, color="666666")


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph_borders(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55 if level == 0 else 1.05)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("•  ")
    p.add_run(text)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_ticker_section(doc, ticker, title, paragraphs, bullets=None):
    add_heading(doc, f"{ticker}  {title}", 2)
    for text in paragraphs:
        doc.add_paragraph(text)
    for text in bullets or []:
        add_bullet(doc, text)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.25)
section.bottom_margin = Cm(2.05)
section.left_margin = Cm(2.25)
section.right_margin = Cm(2.25)
section.header_distance = Cm(1.0)
section.footer_distance = Cm(1.0)

# Base typography
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(34, 34, 34)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.28

for style_name, size, before, after in (
    ("Title", 24, 0, 12),
    ("Subtitle", 11, 0, 14),
    ("Heading 1", 16, 14, 7),
    ("Heading 2", 12.5, 10, 4),
    ("Heading 3", 11, 7, 3),
):
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = style_name != "Subtitle"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

remove_paragraph_borders(styles["Title"])

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.22

# Header/footer
header_p = section.header.paragraphs[0]
header_p.text = "Candy老师美东9月14日夜盘直播详细总结"
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header_p.runs:
    set_run_font(run, 8.5, color="777777")
add_page_number(section.footer.paragraphs[0])

# Cover/opening
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
remove_paragraph_borders(title)
run = title.add_run("Candy老师美东9月14日夜盘")
set_run_font(run, 22, bold=True, color="000000")
run.add_break()
run = title.add_run("直播详细总结")
set_run_font(run, 22, bold=True, color="000000")

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.add_run("美股 港股与主要标的技术面观点整理")

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(10)
r = intro.add_run("核心结论  ")
set_run_font(r, 11, bold=True, color="8B2E2E")
intro.add_run(
    "Candy老师对短线市场保持谨慎，但并不判断会出现深度下跌。她强调利率决议前后方向可能反复，"
    "应先确认支撑、底分型和向上结构，再参与反弹；临近压力位不要追涨，日线下行阶段尽量避免隔夜持有期权。"
)

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta_data = [
    ("直播标题", "美东9.14夜盘 美股 港股分享 MU SOXL AAPL MSTR COIN 恒生指数等怎么看"),
    ("频道", "Candy美股"),
    ("时长", "2小时12分37秒"),
    ("原始视频", VIDEO_URL),
]
for i, (label, value) in enumerate(meta_data):
    c0, c1 = meta.rows[i].cells
    set_cell_width(c0, 3.0)
    set_cell_width(c1, 13.0)
    c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(c0)
    set_cell_margins(c1)
    set_cell_shading(c0, "E9EEF3")
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p0.add_run(label)
    set_run_font(rr, 9.5, bold=True)
    p1 = c1.paragraphs[0]
    if label == "原始视频":
        add_hyperlink(p1, "打开YouTube直播回放", value)
    else:
        rr = p1.add_run(value)
        set_run_font(rr, 9.5)
set_table_borders(meta)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(8)
rr = note.add_run("说明：")
set_run_font(rr, 9, bold=True, color="555555")
rr = note.add_run("以下内容为直播观点整理，不构成投资建议。具体点位会随均线、K线和盘面变化而变化。")
set_run_font(rr, 9, color="555555")

add_heading(doc, "一 直播核心框架", 1)
for text in [
    "先判断日线处于向上笔、向下笔还是震荡，再到30分钟、15分钟或3分钟级别寻找支撑与结构。",
    "价格下跌本身不是买点。Candy老师反复要求同时看到支撑、底分型，以及新的向上中枢或向上离开段。",
    "股票可以在有效支撑附近分批参与，并用关键位管理风险；期权在日线下行时更适合日内操作，不宜随意隔夜。",
    "接近前高、均线密集区或中枢上沿时，不宜追涨。若价格未突破压力，应优先等待回落或确认。",
    "利率决议前后可能出现快速反复。Candy老师个人倾向即使加息也更可能是25个基点，但明确把这视为个人判断。",
]:
    add_numbered(doc, text)

add_heading(doc, "二 市场整体判断", 1)
doc.add_paragraph(
    "Candy老师认为纳斯达克与QQQ的日线结构更接近震荡修复，而不是已经恢复流畅上升。上方均线密集，"
    "她称之为类似“死蜘蛛”的压力结构，短线仍可能回踩。她同时认为夜盘即使下跌，幅度未必很深，"
    "因此不主张在低位盲目追空，也不主张在缺少底部确认时抢反弹。"
)

market_table = doc.add_table(rows=1, cols=4)
market_table.alignment = WD_TABLE_ALIGNMENT.CENTER
market_table.autofit = False
headers = ["市场或主题", "Candy老师的倾向", "关键观察", "更合适的行动"]
widths = [3.2, 3.4, 4.6, 5.0]
for i, h in enumerate(headers):
    cell = market_table.rows[0].cells[i]
    set_cell_width(cell, widths[i])
    set_cell_shading(cell, "17324D")
    set_cell_margins(cell, top=100, bottom=100)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(h)
    set_run_font(rr, 9, bold=True, color="FFFFFF")
set_repeat_table_header(market_table.rows[0])
market_rows = [
    ("纳斯达克 QQQ", "短线谨慎 中期不深度看空", "上方均线与中枢压力 夜盘支撑约25900", "等待小级别底分型和向上结构"),
    ("恒生指数 港股", "不做简单多空判断", "先看趋势线 中枢下沿和底分型", "突破或跌破关键结构后再行动"),
    ("半导体", "优质龙头长期相对有信心", "利率驱动回调与弱势小盘股不同", "股票分批 期权少隔夜 不追压力位"),
    ("黄金 白银", "日线仍在向下笔", "下跌不等于买点 SLV波动更大", "等支撑确认后再做反弹"),
]
for ridx, row_data in enumerate(market_rows):
    cells = market_table.add_row().cells
    for i, value in enumerate(row_data):
        set_cell_width(cells[i], widths[i])
        set_cell_margins(cells[i], top=105, bottom=105)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ridx % 2 == 1:
            set_cell_shading(cells[i], "F2F6F9")
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT
        rr = p.add_run(value)
        set_run_font(rr, 8.7)
set_table_borders(market_table)

add_heading(doc, "三 重点标的速览", 1)
quick = doc.add_table(rows=1, cols=4)
quick.alignment = WD_TABLE_ALIGNMENT.CENTER
quick.autofit = False
headers = ["标的", "短线判断", "直播提到的关键位置", "操作含义"]
widths = [2.4, 3.5, 4.5, 5.8]
for i, h in enumerate(headers):
    cell = quick.rows[0].cells[i]
    set_cell_width(cell, widths[i])
    set_cell_shading(cell, "17324D")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(h)
    set_run_font(rr, 8.7, bold=True, color="FFFFFF")
set_repeat_table_header(quick.rows[0])
quick_rows = [
    ("MU", "反弹中但临近压力", "约940为短线压力", "不追突破 等盘前或开盘确认"),
    ("SOXL", "处于区间中部", "上方约140压力", "等区间低位与底部结构"),
    ("AAPL", "短线偏回调", "331.6偏弱 关注330或更低", "出现新向上中枢后再参与"),
    ("COIN", "低位筹码可观察持有", "160至170为重要成本区", "不追涨 支撑有效才继续持有"),
    ("NVDA", "长期相对有信心", "股票可关注210与200", "接受波动者可分批 否则等日线上拐"),
    ("INTC", "日线支撑尚在", "约94.5为重要支撑", "实体跌破则下行风险增加"),
    ("CRWV", "结构不理想", "70至73更有意义 90附近减仓", "压力位不加仓 跳空后不抢"),
    ("AMZN", "支撑仍在", "约250为风险管理位", "股票可继续持有并观察支撑"),
    ("HOOD", "趋势线之上偏多", "112买入者可参考115附近", "已有仓位可持有或适度减仓"),
    ("PLTR", "高位出现一卖", "等待后续三买结构", "不追高"),
    ("CRWD", "位置偏高", "关注5日均线或15分钟支撑", "等待回踩 不强行买入"),
]
for ridx, row_data in enumerate(quick_rows):
    cells = quick.add_row().cells
    for i, value in enumerate(row_data):
        set_cell_width(cells[i], widths[i])
        set_cell_margins(cells[i], top=90, bottom=90)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ridx % 2 == 1:
            set_cell_shading(cells[i], "F2F6F9")
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 3 else WD_ALIGN_PARAGRAPH.LEFT
        rr = p.add_run(value)
        set_run_font(rr, 8.3)
set_table_borders(quick)

add_heading(doc, "四 个股与板块详细整理", 1)

add_ticker_section(
    doc,
    "恒生指数与港股",
    "等待下跌结构转向",
    [
        "直播开段讨论港股时，Candy老师没有给出笼统的看多或看空结论。她认为，即使价格来到支撑附近，若没有形成底分型，也不能仅凭“跌得多”就买入。",
        "观察重点是趋势线、红色中枢下沿以及小级别结构。若有效跌破这些位置，应先退出或降低仓位；若下跌趋势转为向上，再考虑参与。",
    ],
    [
        "腾讯：具备一定反弹条件，但约450附近有明显压力。近期低位买入者可在压力附近适度减仓，这还不能直接定义为新一轮牛市。",
        "京东9618：日线处于修复或向上笔，只要约104不破可以继续观察或持有，第一道压力约112。",
    ],
)

add_ticker_section(
    doc,
    "纳斯达克与QQQ",
    "震荡修复尚未转强",
    [
        "Candy老师认为日线更像横向整理，上方多条均线聚集，压力尚未消化。夜盘短线可能继续下探，但她不预期会出现特别深的单边下跌。",
        "直播提到纳指约25900附近的短线支撑。不过，当时3分钟级别没有理想的底部结构，因此更合理的做法是等待，而不是提前猜底。",
    ],
)

add_ticker_section(
    doc,
    "NVDA与半导体",
    "长期逻辑与短线节奏分开处理",
    [
        "Candy老师对优质核心科技股的长期表现相对有信心，认为利率因素造成的回调，与基本面较弱的小盘股下跌并不完全相同。但长期看好不代表任何价位都适合买入。",
        "对于NVDA股票，她提到约210和200可作为分批观察区域，前提是投资者能承受继续回撤。若不愿承受这种波动，应等日线重新形成向上笔。日线向下期间，不建议用期权抄底并隔夜持有。",
    ],
    [
        "TSM：热度弱于MU，行情是否启动更依赖盘前和开盘后的确认。",
        "AMD：当时日线形态不理想，Candy老师没有给出积极买入建议。",
        "MSFT：直播中提到此前在490上方已有买入，但没有形成新的独立结论。",
        "QCOM：相对部分同行表现更好，不过结构确认仍不完整，继续等待更合适。",
    ],
)

add_ticker_section(
    doc,
    "MU",
    "有反弹但不适合追高",
    [
        "直播画面中MU约为924，Candy老师把约940视为短线压力。夜盘成交量不足，形态可靠性相对较低，即使向上突破，离开段的空间也可能有限。",
        "她更倾向等盘前或正式开盘后的量价确认，而不是在压力位附近直接追涨。直播中展示的会员盈利案例具有推广性质，不能替代对当前结构的独立判断。",
    ],
)

add_ticker_section(
    doc,
    "SOXL",
    "区间中部缺少性价比",
    [
        "直播结束前，Candy老师认为SOXL处在区间中部，不属于理想进场位置，上方约140存在压力。更好的机会是价格回到区间低位后，出现底分型、新的向上中枢或明确向上离开段。",
        "SOXL属于杠杆ETF。她的框架隐含着更严格的趋势和支撑管理，不适合在缺少确认时长期硬扛。",
    ],
)

add_ticker_section(
    doc,
    "AAPL",
    "短线出现顶部信号",
    [
        "Candy老师认为AAPL盘中出现较清楚的顶分型，可视为短线减仓信号。夜盘小级别走势偏下，并出现较密集的卖出结构。",
        "她认为331.6距离当前价格太近，支撑强度有限，关注点可下移到约330或更低位置。若要重新参与，应等待30分钟级别形成新的向上中枢，或出现更明确的底部确认。",
        "这不是对苹果长期基本面的看空，而是针对当时短线结构提出的回调和消化判断。",
    ],
)

add_ticker_section(
    doc,
    "COIN",
    "低位持仓与新开仓应区别处理",
    [
        "直播画面中COIN约为191.45。Candy老师认为，从较低位置买入的股票仓位，只要关键支撑和趋势线没有破坏，可以继续持有观察。",
        "她没有看到明确的派发结构，后续可能形成收敛或三角整理。约160至170被视为重要成本区。已有低成本仓位与准备新开仓的人应采用不同策略：前者可管理支撑继续持有，后者不宜追涨，应等待支撑和底分型。",
    ],
)

add_ticker_section(
    doc,
    "CRWV",
    "上方压力明显",
    [
        "Candy老师认为CRWV当时的结构不够理想，上方压力较重。对于约93附近被套的仓位，她不建议在压力位继续加仓，而应关注更低的日线支撑。",
        "直播提到约70至73区域更有技术意义；若价格反弹接近90，可考虑减仓或退出。跳空下跌后也不宜立刻抢反弹，应等待结构稳定。",
    ],
)

add_ticker_section(
    doc,
    "INTC",
    "以94.5附近支撑作为风险边界",
    [
        "Candy老师认为INTC日线支撑尚未破坏。若次日高开，可能带动新的向上笔。她把约94.5视为重要支撑，约97附近买入的仓位可以据此进行风险管理。",
        "如果K线实体有效跌破支撑，意味着下行风险增加，不应继续用原来的支撑逻辑解释走势。",
    ],
)

add_ticker_section(
    doc,
    "其他美股",
    "以位置和结构决定操作",
    [],
    [
        "CRCL：当时支撑未破，不宜追空；若要做空，更适合在小级别压力处进行日内交易。",
        "PLTR：价格处于相对高位，并出现“一卖”信号，不适合追涨，应等待后续“三买”结构。",
        "CRWD：位置偏高、下方有效支撑较远，可等待5日均线或15分钟支撑，不要勉强进场。",
        "AMZN：支撑仍然存在，股票仓位可继续持有；约250可作为风险管理参考，约260买入者可以等待结构演化。",
        "HOOD：只要价格维持在两条主要趋势线之上，结构仍偏多。已有股票仓位可以持有；针对约112买入的提问，Candy老师提到约115附近可适度减仓。",
    ],
)

add_ticker_section(
    doc,
    "黄金 白银与AGQ",
    "日线下行时不要把便宜当成买点",
    [
        "Candy老师认为黄金和白银当时都处于日线向下笔。GLD与SLV方向相关，但SLV波动更大。价格已经下跌并不能单独构成反弹买点，仍需支撑与底部结构确认。",
        "对于AGQ，她更倾向在反弹走强时卖出或降低仓位，而不是在日线下行阶段继续加仓。",
    ],
)

add_heading(doc, "五 时间轴与回看入口", 1)
timeline = [
    ("00:04:20", "恒生指数与港股", 260),
    ("00:10:25", "纳斯达克与QQQ", 625),
    ("00:16:25", "NVDA与半导体", 985),
    ("00:55:30", "MU", 3330),
    ("01:10:30", "CRWV", 4230),
    ("01:25:28", "COIN", 5128),
    ("01:36:58", "AAPL", 5818),
    ("02:00:50", "AMZN", 7250),
    ("02:03:45", "京东9618", 7425),
    ("02:08:00", "SOXL", 7680),
]
timeline_table = doc.add_table(rows=1, cols=3)
timeline_table.alignment = WD_TABLE_ALIGNMENT.CENTER
timeline_table.autofit = False
for i, (h, width) in enumerate(zip(("时间", "主题", "回看"), (3.0, 8.0, 5.2))):
    cell = timeline_table.rows[0].cells[i]
    set_cell_width(cell, width)
    set_cell_shading(cell, "17324D")
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(h)
    set_run_font(rr, 9, bold=True, color="FFFFFF")
set_repeat_table_header(timeline_table.rows[0])
for ridx, (ts, topic, seconds) in enumerate(timeline):
    cells = timeline_table.add_row().cells
    for i, width in enumerate((3.0, 8.0, 5.2)):
        set_cell_width(cells[i], width)
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ridx % 2 == 1:
            set_cell_shading(cells[i], "F2F6F9")
    p = cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(ts)
    set_run_font(rr, 9)
    p = cells[1].paragraphs[0]
    rr = p.add_run(topic)
    set_run_font(rr, 9)
    p = cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(p, "跳转观看", f"{VIDEO_URL}&t={seconds}s")
set_table_borders(timeline_table)

add_heading(doc, "六 信息边界与阅读提示", 1)
doc.add_paragraph(
    "直播中多次介绍TodayChart、会员服务、联系方式及会员盈利截图，部分所谓“机构成本”或具体买点仅向会员展示。"
    "这些内容带有明显推广属性，阅读时应与公开技术分析观点分开看待。"
)
doc.add_paragraph(
    "标题中包含MSTR，但在可核对的完整内容与关键画面中，没有形成足够清晰、独立且可复述的MSTR结论。"
    "因此本文不补写或推断其观点。"
)
doc.add_paragraph(
    "最值得保留的操作纪律是：日线下行时不猜底，压力位附近不追涨；股票可围绕支撑分批管理，"
    "期权尽量控制在日内；只有当底分型和向上结构出现后，反弹交易才有更好的依据。"
)

# Prevent widows/orphans and normalize all runs.
for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True
    for run in paragraph.runs:
        if run.font.name is None:
            set_run_font(run)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.widow_control = True
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    if run.font.name is None:
                        set_run_font(run, 9)

doc.core_properties.title = "Candy老师美东9月14日夜盘直播详细总结"
doc.core_properties.subject = "美股 港股 MU SOXL AAPL COIN等直播观点整理"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.keywords = "Candy美股, 美东夜盘, MU, SOXL, AAPL, COIN, 恒生指数"
doc.save(OUTPUT)
print(OUTPUT)
